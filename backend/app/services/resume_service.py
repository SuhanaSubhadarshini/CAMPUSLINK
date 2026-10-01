import io
import logging
import os
import zipfile
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

from ..database import BASE_DIR
from .skill_extractor import extract_skills

UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", str(BASE_DIR / "uploads"))).resolve()
MAX_BYTES = 5 * 1024 * 1024


def file_path(name):
    path = (UPLOAD_DIR / name).resolve()
    if path.parent != UPLOAD_DIR:
        raise HTTPException(400, "Invalid file path")
    return path


def remove_file(name):
    try:
        file_path(name).unlink(missing_ok=True)
    except OSError:
        logging.getLogger(__name__).warning("Could not clean up stored resume")


def extract_text(data, suffix):
    try:
        if suffix == ".txt":
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Binary content")
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("Invalid PDF signature")
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted or len(reader.pages) > 30:
                raise ValueError("Encrypted or too many pages")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("Expanded DOCX is too large")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("Not a DOCX")
            document = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        return text[:200_000]
    except Exception as exc:
        raise HTTPException(422, "Unable to read resume. Use UTF-8 TXT, unencrypted PDF (up to 30 pages), or valid DOCX.") from exc


def save_resume(upload):
    filename = (upload.filename or "resume").replace("\\", "/").split("/")[-1][:200]
    suffix = Path(filename).suffix.lower()
    try:
        if suffix not in {".pdf", ".docx", ".txt"}:
            raise HTTPException(415, "Supported resume formats: PDF, DOCX, TXT")
        data = upload.file.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:
            raise HTTPException(413, "Resume exceeds the 5 MiB limit")
        if not data:
            raise HTTPException(422, "Resume is empty")
        text = extract_text(data, suffix)
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        storage_name = uuid4().hex + suffix
        file_path(storage_name).write_bytes(data)
        return {"filename": filename, "storage_name": storage_name, "text": text, "skills": extract_skills(text)}
    finally:
        upload.file.close()
